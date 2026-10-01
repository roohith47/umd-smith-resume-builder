#!/usr/bin/env python3
"""Build a one-page resume .docx from a structured JSON description.

House visual style (see references/visual-spec.md for the full rationale):
  - Times New Roman throughout. Name 14pt, section headers 11pt, body 10pt.
  - 0.5in margins on all four sides (Word's "Narrow" preset).
  - Section headers: bold, ALL CAPS, left-aligned, with a bottom border rule.
  - Two-column header/sub-header lines (bold org/location, italic title/dates)
    using a right tab stop, matching the candidate's own resume layout.
  - Bullets are plain hyphens, left-aligned, with justified body text.

Usage:
    python3 build_resume.py <input.json> <output.docx>

See references/resume-schema.md for the JSON schema, or ../assets/resume.template.json
for a blank placeholder structure to copy for a new user.
"""

import json
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401  (kept for future table-based layouts)
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "Times New Roman"
MARGIN_IN = 0.5
NAME_SIZE = Pt(14)
HEADER_SIZE = Pt(11)
BODY_SIZE = Pt(10)
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0


def set_run_font(run, size=BODY_SIZE, bold=False, italic=False, name=FONT_NAME):
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    # Make sure complex-script / east-asian fallback doesn't silently override the font.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def add_bottom_border(paragraph):
    """Add a thin bottom border under a paragraph (used for section headers)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_borders.append(bottom)
    p_pr.append(p_borders)


def usable_width_in(doc):
    section = doc.sections[0]
    return (section.page_width - section.left_margin - section.right_margin) / 914400


def set_margins(doc):
    for section in doc.sections:
        section.page_width = Inches(PAGE_WIDTH_IN)
        section.page_height = Inches(PAGE_HEIGHT_IN)
        section.top_margin = Inches(MARGIN_IN)
        section.bottom_margin = Inches(MARGIN_IN)
        section.left_margin = Inches(MARGIN_IN)
        section.right_margin = Inches(MARGIN_IN)
        section.header_distance = Inches(0)
        section.footer_distance = Inches(0)


def add_two_col_paragraph(doc, left_text, right_text, bold=False, italic=False,
                           size=BODY_SIZE, space_after=0):
    """A paragraph with left-aligned text and right-aligned text on the same line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    width = usable_width_in(doc)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
    run_left = p.add_run(left_text)
    set_run_font(run_left, size=size, bold=bold, italic=italic)
    if right_text:
        run_tab = p.add_run("\t")
        set_run_font(run_tab, size=size, bold=bold, italic=italic)
        run_right = p.add_run(right_text)
        set_run_font(run_right, size=size, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("-\t" + text)
    set_run_font(run, size=BODY_SIZE)
    return p


def add_name_contact(doc, data):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(data["name"].upper())
    set_run_font(run, size=NAME_SIZE, bold=True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0 if data.get("work_auth_line") else 8)
    run2 = p2.add_run(data["contact_line"])
    set_run_font(run2, size=BODY_SIZE)

    if data.get("work_auth_line"):
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(8)
        run3 = p3.add_run(data["work_auth_line"])
        set_run_font(run3, size=BODY_SIZE)


def add_section_header(doc, heading):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(heading.upper())
    set_run_font(run, size=HEADER_SIZE, bold=True)
    add_bottom_border(p)
    return p


def add_entries_section(doc, section):
    for entry in section.get("entries", []):
        has_header = bool(entry.get("header_left"))
        if has_header:
            add_two_col_paragraph(
                doc, entry["header_left"], entry.get("header_right", ""),
                bold=True, space_after=0,
            )
        sub_blocks = entry.get("sub_blocks", [])
        for i, sub in enumerate(sub_blocks):
            # Safety net: a schema/content mistake can leave header_left empty (e.g. a
            # standalone project with its title accidentally placed in sub_blocks instead
            # of header_left -- see references/resume-schema.md). Rather than silently
            # rendering what should be a bold title in italic, promote the first sub_block
            # to the bold header style so the title still reads correctly.
            is_title_line = not has_header and i == 0
            add_two_col_paragraph(
                doc, sub.get("sub_left", ""), sub.get("sub_right", ""),
                bold=is_title_line, italic=not is_title_line, space_after=0,
            )
            bullets = sub.get("bullets", [])
            for j, bullet in enumerate(bullets):
                add_bullet(doc, bullet, space_after=2 if j == len(bullets) - 1 else 0)


def add_skills_section(doc, section):
    for line in section.get("lines", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run_label = p.add_run(line["label"] + ": ")
        set_run_font(run_label, bold=True)
        run_items = p.add_run(line["items"])
        set_run_font(run_items)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def add_bullets_section(doc, section):
    bullets = section.get("bullets", [])
    for i, bullet in enumerate(bullets):
        add_bullet(doc, bullet, space_after=2 if i == len(bullets) - 1 else 0)


SECTION_BUILDERS = {
    "entries": add_entries_section,
    "skills": add_skills_section,
    "bullets": add_bullets_section,
}


def build(data, output_path):
    doc = Document()
    # Base style, so any paragraph that falls back to Normal still matches.
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    set_margins(doc)

    add_name_contact(doc, data)

    for section in data.get("sections", []):
        builder = SECTION_BUILDERS.get(section["type"])
        if builder is None:
            raise ValueError(f"Unknown section type: {section['type']!r}")
        add_section_header(doc, section["heading"])
        builder(doc, section)

    doc.save(output_path)
    return doc


def estimate_lines(doc):
    """Rough line-count estimate for a one-page sanity check (heuristic, not exact).

    Times New Roman at 10pt averages roughly 15.5 characters per inch of line width.
    This undercounts slightly for bold/italic runs and overcounts for short lines, so
    treat the result as a ballpark, not a guarantee -- always eyeball the actual file.
    """
    width_in = usable_width_in(doc)
    chars_per_line = {10: 15.5, 11: 14.0, 14: 11.0}
    total_lines = 0.0
    for p in doc.paragraphs:
        text = "".join(r.text for r in p.runs).replace("\t", "   ")
        if not text.strip():
            total_lines += 0.3
            continue
        size_pt = 10
        if p.runs:
            size_pt = int((p.runs[0].font.size or Pt(10)).pt)
        cpl = chars_per_line.get(size_pt, 15.5) * width_in
        total_lines += max(1, -(-len(text) // max(1, int(cpl))))  # ceil div
    return total_lines


def estimate_fits_one_page(doc):
    lines = estimate_lines(doc)
    # ~11pt line height at 10pt body on a 10in usable height (11in page - 1in margins).
    usable_height_in = PAGE_HEIGHT_IN - 2 * MARGIN_IN
    max_lines = usable_height_in * 72 / 12.5  # ~12.5pt per line incl. paragraph spacing
    return lines, max_lines, lines <= max_lines


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    input_path, output_path = sys.argv[1], sys.argv[2]
    with open(input_path) as f:
        data = json.load(f)
    doc = build(data, output_path)
    lines, max_lines, fits = estimate_fits_one_page(doc)
    print(f"Wrote {output_path}")
    print(f"Estimated content: ~{lines:.0f} lines (budget ~{max_lines:.0f} for one page).")
    if not fits:
        print("WARNING: this draft is estimated to run past one page. "
              "Trim a bullet or tighten wording -- don't shrink below the house font/margins. "
              "This is a heuristic: open the file to confirm visually either way.")
    else:
        print("Looks like it should fit one page -- open the file to confirm visually.")


if __name__ == "__main__":
    main()
